"""Genera la matriz de competencias en LaTeX (tabularray) y en Excel desde los mismos datos."""
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent

INSTR = ["Examen escrito", "Cuestionario (GeoGebra Classroom)", "Fichas de construcción en GeoGebra",
         "Libro de GeoGebra (trabajo en equipo)", "Exposición oral", "Mapa conceptual",
         "Observación sistemática (registro y Kanban)", "Coevaluación y autoevaluación (rúbrica)"]
CORTO = ["Examen", "Cuestionario", "Fichas GeoGebra", "Libro GeoGebra", "Exposición oral",
         "Mapa conceptual", "Observación", "Co/autoevaluación"]
EX, CU, FI, LI, EP, MC, OB, CO = range(8)

CRIT = {
    "1.1": "Manejar algunas estrategias y herramientas, incluidas las digitales, en la modelización y resolución de problemas de la vida cotidiana y de la ciencia y la tecnología, evaluando su eficiencia en cada caso.",
    "1.2": "Obtener todas las posibles soluciones matemáticas de problemas de la vida cotidiana y de la ciencia y la tecnología, describiendo el procedimiento utilizado.",
    "2.1": "Comprobar la validez matemática de las posibles soluciones de un problema utilizando el razonamiento y la argumentación.",
    "2.2": "Seleccionar la solución más adecuada de un problema en función del contexto, usando el razonamiento y la argumentación.",
    "3.1": "Adquirir nuevo conocimiento matemático a partir de la formulación de conjeturas y problemas de forma guiada.",
    "3.2": "Emplear herramientas tecnológicas adecuadas en la formulación o investigación de conjeturas o problemas.",
    "5.1": "Manifestar una visión matemática integrada, investigando y conectando las diferentes ideas matemáticas.",
    "5.2": "Resolver problemas en contextos matemáticos estableciendo y aplicando conexiones entre las diferentes ideas matemáticas.",
    "7.1": "Representar ideas matemáticas, estructurando diferentes razonamientos matemáticos y seleccionando las tecnologías más adecuadas.",
    "7.2": "Seleccionar y utilizar diversas formas de representación, valorando su utilidad para compartir información.",
    "8.1": "Mostrar organización al comunicar las ideas matemáticas empleando el soporte, la terminología y el rigor apropiados.",
    "9.1": "Afrontar las situaciones de incertidumbre, aceptando y aprendiendo del error como parte del proceso de aprendizaje de las matemáticas.",
    "9.2": "Mostrar una actitud positiva y perseverante, aceptando y aprendiendo de la crítica razonada al hacer frente a las diferentes actividades de las matemáticas.",
    "9.3": "Participar en tareas matemáticas de forma activa en equipos heterogéneos, escuchando su razonamiento.",
}

# (bloque, [contenidos literales del Decreto 64/2022], [(criterio, operativizado, instrumentos)])
GRUPOS = [
    ("A. Números y operaciones: vectores", [
        "Adición y producto escalar de vectores: propiedades y representaciones.",
        "Conjunto de vectores: estructura, comprensión y propiedades.",
        "Módulo de un vector, coordenadas de un vector con respecto a una base, ángulo entre dos vectores y proyección ortogonal.",
        "Estrategias para operar con números reales y vectores: cálculo mental o escrito en los casos sencillos y con herramientas tecnológicas en los casos más complicados.",
    ], [
        ("7.1", "Representar vectores en GeoGebra y operar con ellos (suma, producto por un número y producto escalar), interpretando geométricamente el resultado de cada operación.", [EX, CU, FI]),
        ("2.1", "Comprobar con GeoGebra los cálculos hechos a mano del módulo, el ángulo entre dos vectores y la proyección ortogonal, argumentando el origen de las discrepancias.", [EX, FI]),
        ("5.2", "Resolver problemas de perpendicularidad y de cálculo de ángulos mediante el producto escalar, conectando su expresión en coordenadas con su significado geométrico.", [EX, CU]),
    ]),
    ("B y D. Trigonometría", [
        "Razones trigonométricas de un ángulo cualquiera medido en grados o en radianes.",
        "Demostración de las identidades trigonométricas. Razones trigonométricas del ángulo suma, el ángulo diferencia, el ángulo doble y el ángulo mitad.",
        "Cálculo de las razones trigonométricas de un ángulo cualquiera empleando las principales fórmulas trigonométricas.",
        "Conjeturas geométricas en el plano: validación por medio de la deducción y la demostración de teoremas.",
        "Resolución de ecuaciones trigonométricas sencillas.",
    ], [
        ("3.2", "Emplear una circunferencia goniométrica dinámica en GeoGebra (con un deslizador para el ángulo) para investigar cómo varían el seno, el coseno y la tangente en los cuatro cuadrantes.", [FI, OB]),
        ("3.1", "Formular, de forma guiada, conjeturas sobre identidades trigonométricas (fundamental, ángulo doble, ángulo suma) a partir de la construcción y demostrarlas después por escrito.", [EX, FI]),
        ("1.2", "Obtener todas las soluciones de ecuaciones trigonométricas sencillas, tanto en [0°, 360°) como su expresión general, describiendo el procedimiento seguido.", [EX, CU]),
    ]),
    ("B. Medida y geometría: resolución de triángulos", [
        "Cálculo de longitudes y medidas angulares: uso de la trigonometría.",
        "Aplicación de las razones trigonométricas, el teorema de los senos y el teorema del coseno en la resolución de triángulos y de problemas geométricos de contexto real.",
        "Demostración del teorema del seno y del coseno.",
    ], [
        ("1.1", "Elegir entre el teorema del seno y el del coseno según los datos de cada triángulo, valorando qué estrategia resulta más eficiente.", [EX, FI]),
        ("1.2", "Obtener todas las soluciones de un triángulo dado por dos lados y el ángulo opuesto a uno de ellos (dos, una o ninguna), describiendo el procedimiento utilizado.", [EX, FI]),
        ("2.1", "Verificar en GeoGebra, arrastrando los vértices, la validez de los triángulos obtenidos analíticamente y argumentar por qué un caso no tiene solución.", [FI, CU]),
    ]),
    ("C. Geometría en el plano: rectas", [
        "Ecuaciones de la recta en el espacio bidimensional.",
        "Estudio de la posición relativa de puntos y rectas en el plano.",
        "Planteamiento y resolución de problemas de geometría afín relacionados con la incidencia, el paralelismo y la ortogonalidad de rectas en el plano.",
        "Planteamiento y resolución de problemas de geometría métrica relacionados con la medida de ángulos entre rectas y la medida de distancias entre puntos y rectas.",
        "Expresiones algebraicas de objetos geométricos: selección de la más adecuada en función de la situación a resolver.",
    ], [
        ("5.1", "Relacionar el vector director y el vector normal con las distintas ecuaciones de la recta, integrando vectores, trigonometría y geometría analítica en una visión de conjunto.", [MC, EX]),
        ("5.2", "Estudiar la posición relativa de dos rectas y calcular ángulos entre rectas y distancias de un punto a una recta, aplicando el producto escalar.", [EX, CU]),
        ("7.2", "Seleccionar la ecuación de la recta (vectorial, paramétrica, continua, general o explícita) más útil para cada problema, justificando la elección.", [EX, FI]),
    ]),
    ("C y D. Modelización: producto final «Medir lo que no se puede alcanzar»", [
        "Representación de objetos geométricos en el plano mediante herramientas digitales.",
        "Modelos matemáticos (geométricos, algebraicos, grafos...) en la resolución de problemas en el plano. Conexiones con otras disciplinas y áreas de interés.",
        "Modelización de la posición y el movimiento de un objeto en el plano mediante vectores.",
        "Formulación, resolución y análisis de problemas de la vida cotidiana y de la ciencia y la tecnología utilizando herramientas o programas adecuados.",
    ], [
        ("1.1", "Diseñar un método de medida y un modelo a escala en GeoGebra para calcular una altura o distancia inaccesible del centro, evaluando su eficiencia frente a otros métodos.", [LI, EP]),
        ("2.2", "Seleccionar el resultado más razonable de entre las medidas repetidas, valorando el efecto que tiene un error de un grado en el resultado final.", [LI, EP]),
        ("7.2", "Organizar en un libro de GeoGebra la construcción, los cálculos y la explicación del método, eligiendo las representaciones más claras para el resto de la clase.", [LI, CO]),
        ("8.1", "Exponer oralmente el método y los resultados con orden, usando la terminología y la notación matemáticas con rigor.", [EP, CO]),
    ]),
    ("F. Actitudes y aprendizaje", [
        "Tratamiento del error como elemento movilizador de conocimientos previos adquiridos y generador de oportunidades de aprendizaje en el aula de matemáticas.",
        "Reconocimiento y aceptación de diversos planteamientos en la resolución de problemas, transformando los enfoques de los demás en nuevas y mejoradas estrategias propias.",
        "Técnicas y estrategias de trabajo en equipo para la resolución de problemas y tareas matemáticas, en grupos heterogéneos.",
        "Destrezas para desarrollar una comunicación efectiva: la escucha activa, la formulación de preguntas o solicitud y prestación de ayuda cuando sea necesario.",
    ], [
        ("9.1", "Aceptar las conjeturas fallidas y los errores de medida como punto de partida para revisar el razonamiento en la puesta en común.", [OB, CO]),
        ("9.2", "Incorporar la crítica razonada recibida en la coevaluación para mejorar el libro de GeoGebra, sin abandonar ante las dificultades.", [CO, OB]),
        ("9.3", "Participar activamente en el equipo heterogéneo, asumiendo tareas en el tablero Kanban y escuchando los razonamientos de los compañeros.", [OB, CO]),
    ]),
]


def tex():
    n = len(INSTR)
    rot = lambda s: r"\rotatebox{90}{" + s + "}"
    L = [r"\begin{longtblr}[caption={Matriz de competencias de la situación de aprendizaje}, label={tab:matriz}]{",
         r"  colspec={Q[l,wd=7.0cm] Q[l,wd=5.0cm] Q[l,wd=7.2cm] *{%d}{Q[c,m,wd=0.55cm]}}," % n,
         r"  hlines={0.4pt,gris!50}, vlines={0.4pt,gris!50}, rowhead=2, colsep=3pt,",
         r"  row{1-2}={font=\bfseries, fg=white, bg=acento, m}, cells={font=\scriptsize},",
         r"}",
         r"\SetCell[r=2]{c} Contenidos (Decreto 64/2022) & \SetCell[r=2]{c} Criterio de evaluación & "
         r"\SetCell[r=2]{c} Criterio operativizado & \SetCell[c=%d]{c} Instrumentos de evaluación" % n + " &" * (n - 1) + r" \\",
         " & & & " + " & ".join(rot(s) for s in CORTO) + r" \\"]
    for bloque, cont, filas in GRUPOS:
        celda = r"{\textbf{" + bloque + r"}\\ " + r"\\ ".join(r"\textbullet~" + c for c in cont) + "}"
        for i, (c, op, ins) in enumerate(filas):
            first = r"\SetCell[r=%d]{l,m} %s" % (len(filas), celda) if i == 0 else ""
            marks = " & ".join(r"\textbf{X}" if k in ins else "" for k in range(n))
            end = r"\\" if i == len(filas) - 1 else r"\\*"  # ponytail: no partir un grupo entre páginas
            L.append(rf"{first} & \textbf{{{c}.}} {CRIT[c]} & {op} & {marks} {end}")
    L.append(r"\end{longtblr}")
    (OUT / "matriz_tabla.tex").write_text("\n".join(L) + "\n", encoding="utf-8")


def xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz de competencias"
    head = PatternFill("solid", fgColor="B00020")
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    for col, t in enumerate(["Contenidos (Decreto 64/2022)", "Criterio de evaluación", "Criterio operativizado"], 1):
        ws.cell(1, col, t)
        ws.merge_cells(start_row=1, end_row=2, start_column=col, end_column=col)
    ws.cell(1, 4, "Instrumentos de evaluación")
    ws.merge_cells(start_row=1, end_row=1, start_column=4, end_column=3 + len(INSTR))
    for k, s in enumerate(INSTR):
        ws.cell(2, 4 + k, s).alignment = Alignment(text_rotation=90, wrap_text=True, horizontal="center", vertical="bottom")
    r = 3
    for bloque, cont, filas in GRUPOS:
        ws.cell(r, 1, bloque + "\n" + "\n".join("• " + c for c in cont))
        ws.merge_cells(start_row=r, end_row=r + len(filas) - 1, start_column=1, end_column=1)
        for c, op, ins in filas:
            ws.cell(r, 2, f"{c}. {CRIT[c]}")
            ws.cell(r, 3, op)
            for k in ins:
                ws.cell(r, 4 + k, "X")
            ws.row_dimensions[r].height = 95
            r += 1
    for row in ws.iter_rows(min_row=1, max_row=r - 1, max_col=3 + len(INSTR)):
        for cell in row:
            cell.border = box
            if cell.row <= 2:
                cell.fill, cell.font = head, Font(bold=True, color="FFFFFF")
                if cell.row == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif cell.column <= 3:
                cell.alignment = wrap
            else:
                cell.alignment, cell.font = Alignment(horizontal="center", vertical="center"), Font(bold=True)
    ws.row_dimensions[2].height = 190
    for col, w in zip("ABC", (55, 40, 55)):
        ws.column_dimensions[col].width = w
    for k in range(len(INSTR)):
        ws.column_dimensions[chr(ord("D") + k)].width = 6
    ws.freeze_panes = "D3"
    wb.save(OUT / "Matriz_competencias_DCM.xlsx")


assert all(c in CRIT for _, _, f in GRUPOS for c, _, _ in f)
assert set(range(len(INSTR))) == {k for _, _, f in GRUPOS for *_, ins in f for k in ins}, "instrumento sin usar"
tex()
xlsx()
